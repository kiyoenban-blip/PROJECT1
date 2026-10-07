import asyncio
import hashlib
import json
import os
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook
from frr.engine import inspect_report_package, recheck_corrected_report
from frr.workbook import read_review
from frr.server import mcp

class ReviewTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(); self.root=Path(self.temp.name)
        self.cwd=Path.cwd(); os.chdir(self.root)
    def tearDown(self):
        os.chdir(self.cwd); self.temp.cleanup()
    def facts(self,name,value=52430,**kw):
        p=self.root/name
        row=dict(concept='Revenue',value=value,period='2026',unit='KRW',entity='ABC',scope='consolidated',location='Revenue')
        row.update(kw); p.write_text(json.dumps({'facts':[row]})); return str(p)
    def test_reconcile_materiality_and_immutable(self):
        p=self.facts('primary.json'); r=self.facts('reference.json',52340)
        before=Path(p).read_bytes()
        result=inspect_report_package(p,[r],materiality=500,performance_materiality=350,trivial_threshold=25)
        f=result['findings'][0]
        self.assertEqual(f['difference_numeric'],'90')
        self.assertEqual(f['materiality_ratio'],'0.18')
        self.assertEqual(before,Path(p).read_bytes())
        self.assertEqual(read_review(result['output_file'])['findings'][0]['issue_id'],f['issue_id'])
        self.assertEqual(result['final_status'],'REVIEW_COMPLETE_WITH_OPEN_ITEMS')
    def test_clearing_and_manual_fields(self):
        p=self.facts('primary.json'); ref=self.facts('ref.json',52340)
        first=inspect_report_package(p,[ref])
        wb=load_workbook(first['output_file']); wb['02_수정사항']['R2']='Kim'; wb.save(first['output_file']); wb.close()
        second=recheck_corrected_report(first['output_file'],p)
        self.assertEqual(second['summary']['UNRESOLVED'],1)
        self.assertEqual(second['findings'][0]['owner'],'Kim')
        fixed=self.facts('fixed.json',52340)
        third=recheck_corrected_report(second['output_file'],fixed)
        self.assertEqual(third['summary']['RESOLVED'],1)
        self.assertEqual(third['findings'][0]['issue_id'],first['findings'][0]['issue_id'])
        fourth=recheck_corrected_report(first['output_file'],fixed,[])
        self.assertEqual(fourth['summary']['NOT_VERIFIABLE'],1)
    def test_metadata_mismatch_does_not_compare(self):
        for override in [dict(unit='USD'),dict(period='2025'),dict(scope='separate'),dict(entity='XYZ'),dict(dimensions='member:A')]:
            p=self.facts('p.json'); r=self.facts('r.json',1,**override)
            result=inspect_report_package(p,[r],export_excel=False)
            self.assertEqual(result['findings'],[])
            self.assertTrue(any(x['status']=='NOT_VERIFIABLE' for x in result['logs']))
    def test_overwrite_guard(self):
        p=self.root/'source.xlsx'; wb=Workbook(); wb.save(p)
        before=p.read_bytes()
        with self.assertRaises(ValueError):
            inspect_report_package(str(p),output_path=str(p))
        self.assertEqual(p.read_bytes(),before)
    def test_excel_rules_and_formula_injection(self):
        p=self.root/'source.xlsx'; wb=Workbook(); ws=wb.active
        ws['H20']='=SUM(H10:H19)'; ws['I20']='=SUM(I10:I19)'; ws['J20']=52340; ws['K20']='=SUM(K10:K19)'
        ws['A1']="='[2025_TB.xlsx]TB'!H150"; ws['A2']='#DIV/0!'; wb.save(p)
        result=inspect_report_package(str(p))
        rules={f['rule_id'] for f in result['findings']}
        self.assertTrue({'HARD_CODE','FORMULA_ERROR','EXTERNAL_LINK'}.issubset(rules))
        wb=load_workbook(result['output_file'])
        self.assertFalse(any(c.data_type=='f' for ws in wb for row in ws for c in row))
        wb.close()
    def test_footing(self):
        p=self.root/'p.json'
        common=dict(period='2026',unit='KRW',entity='ABC',scope='separate')
        p.write_text(json.dumps({'facts':[dict(common,concept=k,location=k,value=v) for k,v in [('a',10),('b',20),('total',31)]],
            'calculations':[{'id':'assets','total':'total','components':['a','b']}]}))
        result=inspect_report_package(str(p),export_excel=False)
        f=next(f for f in result['findings'] if f['rule_id']=='FOOTING')
        self.assertEqual(f['difference'],'1')
    def test_invalid_threshold(self):
        p=self.facts('p.json')
        for n in [0,-1,float('nan'),float('inf')]:
            with self.assertRaises(ValueError): inspect_report_package(p,materiality=n)
    def test_xbrl(self):
        p=self.root/'report.xbrl'
        p.write_text('''<xbrl xmlns="http://www.xbrl.org/2003/instance" xmlns:t="urn:test"><context id="c"><entity><identifier scheme="id">ABC</identifier></entity><period><instant>2026-12-31</instant></period></context><unit id="u"><measure>KRW</measure></unit><t:Revenue contextRef="c" unitRef="u" decimals="bad">52340</t:Revenue><t:Assets contextRef="missing" unitRef="u">10</t:Assets></xbrl>''')
        result=inspect_report_package(str(p),export_excel=False,review_config={'xbrl_validation':'native'})
        self.assertEqual({f['rule_id'] for f in result['findings']},{'XBRL_REFERENCE','XBRL_DECIMALS'})
    def test_missing_formula_cache_cannot_clear_error(self):
        p=self.root/'broken.xlsx'; wb=Workbook(); wb.active['A1']='#DIV/0!'; wb.save(p)
        first=inspect_report_package(str(p))
        fixed=self.root/'unverified.xlsx'; wb=Workbook(); wb.active['A1']='=1/0'; wb.save(fixed)
        second=recheck_corrected_report(first['output_file'],str(fixed))
        self.assertEqual(second['summary']['NOT_VERIFIABLE'],1)
        self.assertEqual(second['summary']['RESOLVED'],0)
    def test_recurrence_keeps_issue_id(self):
        p=self.facts('primary.json'); ref=self.facts('ref.json',52340)
        first=inspect_report_package(p,[ref])
        second=recheck_corrected_report(first['output_file'],self.facts('fixed.json',52340))
        third=recheck_corrected_report(second['output_file'],p)
        self.assertEqual(third['summary']['UNRESOLVED'],1)
        self.assertEqual(first['findings'][0]['issue_id'],third['findings'][0]['issue_id'])
    def test_existing_output_preserved(self):
        p=self.facts('p.json'); out=self.root/'existing.xlsx'; out.write_bytes(b'preserve')
        with self.assertRaises(FileExistsError):
            inspect_report_package(p,output_path=str(out))
        self.assertEqual(out.read_bytes(),b'preserve')

    def test_mcp_tools(self):
        tools=asyncio.run(mcp.list_tools())
        self.assertEqual({t.name for t in tools},{'inspect_report_package','recheck_corrected_report','compare_report_versions'})

if __name__=='__main__': unittest.main()
