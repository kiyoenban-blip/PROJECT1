import importlib.util
import json
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from openpyxl import Workbook,load_workbook
from frr import inspect_report_package,recheck_corrected_report,compare_report_versions
from frr.workbook import read_review
from frr.parsers import parse
from frr.config import ReviewConfig

class ExtendedTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
    def tearDown(self):
        self.temp.cleanup()
    def write(self,name,data):
        p=self.root/name;p.write_text(json.dumps(data,ensure_ascii=False) if not isinstance(data,str) else data);return str(p)
    def fact(self,value=10,**extra):
        return dict(dict(concept='Revenue',value=str(value),period='2026',unit='KRW',entity='ABC',scope='separate',location='revenue'),**extra)
    def review(self,name='p.json',data=None,refs=None,config=None,**kwargs):
        p=self.write(name,data if data is not None else {'facts':[self.fact()]})
        return inspect_report_package(p,refs,review_config=config,output_path=str(self.root/(name+'.xlsx')),**kwargs)
    def test_all_four_statuses(self):
        ref=self.write('r.json',{'facts':[self.fact()]})
        self.assertEqual(self.review(refs=[ref])['final_status'],'REVIEW_COMPLETE')
        self.assertEqual(self.review('bad.json',{'facts':[self.fact(11)]},[ref])['final_status'],'REVIEW_COMPLETE_WITH_OPEN_ITEMS')
        self.assertEqual(self.review('no-ref.json')['final_status'],'REVIEW_REQUIRED')
        result=self.review('corrupt.xml','<broken')
        self.assertEqual(result['final_status'],'VALIDATION_FAILED')
        self.assertEqual(read_review(result['output_file'])['final_status'],'VALIDATION_FAILED')
    def test_failed_recheck_never_clears(self):
        ref=self.write('r.json',{'facts':[self.fact()]})
        first=self.review(data={'facts':[self.fact(11)]},refs=[ref])
        broken=self.write('broken.json','{')
        second=recheck_corrected_report(first['output_file'],broken,output_path=str(self.root/'recheck.xlsx'))
        self.assertEqual(second['summary']['NOT_VERIFIABLE'],1)
        self.assertEqual(second['final_status'],'VALIDATION_FAILED')
    def test_html_table_extraction_and_footing(self):
        html='<html><body><p>재무제표</p><table><tr><th>계정</th><th>2026</th></tr><tr><td>현금</td><td>10</td></tr><tr><td>예금</td><td>20</td></tr><tr><td>합계</td><td>31</td></tr></table></body></html>'
        result=self.review('table.html',html,config={'metadata':{'entity':'ABC','scope':'separate','unit':'KRW'},'require_references':False})
        self.assertEqual(len(result['facts']),3)
        f=next(f for f in result['findings'] if f['rule_id']=='FOOTING')
        self.assertEqual(f['difference'],'1')
        self.assertEqual(f['confidence'],.8)
    def test_tables_without_metadata_not_guessed(self):
        html='<table><tr><td>계정</td><td>2026</td></tr><tr><td>현금</td><td>10</td></tr></table>'
        result=self.review('table.html',html)
        self.assertEqual(result['facts'],[])
        self.assertEqual(result['final_status'],'REVIEW_REQUIRED')
    def test_dart_dsd_immutable(self):
        p=self.root/'report.dsd'
        with ZipFile(p,'w') as z:
            z.writestr('contents.xml','<DOCUMENT><P>유형 자산</P><TABLE><TR><TH>계정</TH><TH>2026</TH></TR><TR><TD>현금</TD><TD>10</TD></TR></TABLE></DOCUMENT>')
        before=p.read_bytes()
        result=inspect_report_package(str(p),export_excel=False,review_config={'metadata':{'entity':'ABC','unit':'KRW','scope':'separate'}})
        self.assertEqual(before,p.read_bytes())
        self.assertEqual(result['facts'][0]['location'],'contents.xml::table:1:R2C2')
        self.assertTrue(any(f['rule_id']=='TERM' for f in result['findings']))
    def test_archive_traversal_rejected(self):
        p=self.root/'report.zip'
        with ZipFile(p,'w') as z:z.writestr('../escape.txt','not written')
        result=inspect_report_package(str(p),export_excel=False)
        self.assertEqual(result['final_status'],'VALIDATION_FAILED')
        self.assertFalse((self.root.parent/'escape.txt').exists())
    def test_text_expectations_terms_and_notes(self):
        text='회사명: WRONG\n보고기간: 2025.12.31\n유형 자산은 주석 9 참조\n1. 일반사항'
        config={'expectations':[{'id':'company','label':'회사명','selector':'회사명:','expected':'ABC','kind':'company'},
            {'id':'date','label':'보고기간','selector':'보고기간:','expected':'2026-12-31','kind':'date'}]}
        result=self.review('report.txt',text,config=config)
        self.assertEqual({f['rule_id'] for f in result['findings']},{'TERM','EXPECTATION','NOTE_REFERENCE'})
        self.assertEqual(len(result['findings']),4)
        fixed=self.write('fixed.txt','회사명: ABC\n보고기간: 2026.12.31\n유형자산은 주석 9 참조\n1. 일반사항\n9. 유형자산')
        second=recheck_corrected_report(result['output_file'],fixed,output_path=str(self.root/'fixed-review.xlsx'))
        self.assertEqual(second['summary']['RESOLVED'],4)
    def test_reference_excel_issue_clears(self):
        p=self.write('p.json',{'facts':[self.fact()]})
        r=self.root/'ref.xlsx';wb=Workbook();wb.active['A1']='#REF!';wb.save(r)
        first=inspect_report_package(p,[str(r)],output_path=str(self.root/'first.xlsx'))
        wb=Workbook();wb.active['A1']=10;wb.save(r)
        second=recheck_corrected_report(first['output_file'],p,output_path=str(self.root/'second.xlsx'))
        found=next(f for f in second['findings'] if f['rule_id']=='FORMULA_ERROR')
        self.assertEqual(found['clearing'],'RESOLVED')
    def test_dropping_conflicting_reference_not_resolution(self):
        p=self.write('p.json',{'facts':[self.fact()]})
        r1=self.write('r1.json',{'facts':[self.fact()]});r2=self.write('r2.json',{'facts':[self.fact(11)]})
        first=inspect_report_package(p,[r1,r2],output_path=str(self.root/'first.xlsx'))
        second=recheck_corrected_report(first['output_file'],p,[r1],output_path=str(self.root/'second.xlsx'))
        self.assertEqual(second['summary']['NOT_VERIFIABLE'],1)
    def test_duplicate_facts(self):
        result=self.review(data={'facts':[self.fact(),self.fact(11,location='other')]})
        self.assertTrue(any(f['rule_id']=='DUPLICATE_FACT' for f in result['findings']))
        ids=[f['issue_id'] for f in result['findings']]
        self.assertEqual(len(ids),len(set(ids)))
    def test_unit_mapping_materiality(self):
        ref=self.write('r.json',{'facts':[self.fact(1000000)]})
        result=self.review(data={'facts':[self.fact(2,unit='KRW_million')]},refs=[ref],
            config={'unit_mappings':{'KRW_million':{'unit':'KRW','multiplier':'1000000'}},'materiality_unit':'KRW'},materiality=500000)
        f=result['findings'][0]
        self.assertEqual(f['materiality_ratio'],'2')
        self.assertEqual(f['severity'],'CRITICAL')
        self.assertEqual(result['facts'][0]['raw_unit'],'KRW_million')
    def test_mixed_units_do_not_get_wrong_ratios(self):
        ref=self.write('r.json',{'facts':[self.fact(),self.fact(10,concept='Dollar',unit='USD',location='dollar')]})
        result=self.review(data={'facts':[self.fact(11),self.fact(11,concept='Dollar',unit='USD',location='dollar')]},refs=[ref],materiality=5)
        self.assertTrue(all(f['materiality_ratio'] is None for f in result['findings']))
        self.assertEqual(result['final_status'],'REVIEW_REQUIRED')
    def test_weighted_cross_footing(self):
        facts=[self.fact(10,concept='a',location='a'),self.fact(3,concept='b',location='b'),self.fact(8,concept='total',location='total')]
        result=self.review(data={'facts':facts,'calculations':[{'id':'cross','total':'total','components':['a',{'location':'b','weight':-1}],'cross':True}]})
        f=next(f for f in result['findings'] if f['rule_id']=='CROSS_FOOTING')
        self.assertEqual(f['difference'],'1')
    def test_changed_calculation_not_silently_cleared(self):
        facts=[self.fact(10,concept='a',location='a'),self.fact(3,concept='b',location='b'),self.fact(10,concept='total',location='total')]
        first=self.review(data={'facts':facts,'calculations':[{'id':'sum','total':'total','components':['a','b']}]})
        p=self.write('fixed.json',{'facts':facts,'calculations':[{'id':'sum','total':'total','components':['a']}]})
        second=recheck_corrected_report(first['output_file'],p,output_path=str(self.root/'second.xlsx'))
        self.assertEqual(next(f for f in second['findings'] if f['rule_id']=='FOOTING')['clearing'],'NOT_VERIFIABLE')
    def test_version_move_and_unexpected_numeric_change(self):
        old=self.write('old.json',{'facts':[self.fact()], 'texts':{'a':'이동 문단','b':'기준일 2026.12.31'}})
        new=self.write('new.json',{'facts':[self.fact(11)],'texts':{'c':'이동 문단','b':'기준일 2025.12.31'}})
        ref=self.write('r.json',{'facts':[self.fact()]})
        r=compare_report_versions(old,new,[ref],output_path=str(self.root/'review.xlsx'))
        self.assertTrue({'MOVED','DATE_CHANGED','FACT_CHANGED'}.issubset({c['type'] for c in r['version_changes']}))
        self.assertTrue(any(f['rule_id']=='UNEXPECTED_CHANGE' for f in r['findings']))
    def test_prior_comparative_amount(self):
        prior=self.write('prior.json',{'facts':[self.fact(period='2025')]})
        result=self.review(data={'facts':[self.fact(11,period='2025')]},prior_file=prior,config={'require_references':False})
        self.assertTrue(any(f['rule_id']=='PRIOR_PERIOD' for f in result['findings']))
    def test_workbook_metadata_tampering_detected(self):
        r=self.review()
        wb=load_workbook(r['output_file']);ws=wb['_ReviewData'];ws['A1']=ws['A1'].value.replace('"2026"','"2025"');wb.save(r['output_file']);wb.close()
        with self.assertRaises(ValueError):read_review(r['output_file'])
    def test_xml_external_entities_rejected(self):
        result=self.review('evil.xml','<!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><x>&e;</x>')
        self.assertEqual(result['final_status'],'VALIDATION_FAILED')
        self.assertNotIn('root:',json.dumps(result))
    def test_renamed_reference_role_retains_issue(self):
        p=self.write('p.json',{'facts':[self.fact()]})
        r1=self.root/'ref1.xlsx';r2=self.root/'ref2.xlsx'
        wb=Workbook();wb.active['A1']='#REF!';wb.save(r1);wb.save(r2)
        cfg={'reference_roles':{str(r1):'TB',str(r2):'TB'}}
        first=inspect_report_package(p,[str(r1)],review_config=cfg,output_path=str(self.root/'first.xlsx'))
        second=recheck_corrected_report(first['output_file'],p,[str(r2)],output_path=str(self.root/'second.xlsx'))
        self.assertEqual(second['summary']['UNRESOLVED'],1)
    @unittest.skipUnless(importlib.util.find_spec('arelle'),'Arelle dependency not installed')
    def test_arelle_missing_taxonomy_is_not_pass(self):
        xml='''<xbrl xmlns="http://www.xbrl.org/2003/instance" xmlns:t="urn:test" xmlns:link="http://www.xbrl.org/2003/linkbase" xmlns:xlink="http://www.w3.org/1999/xlink"><link:schemaRef xlink:type="simple" xlink:href="missing.xsd"/><t:Revenue contextRef="missing" unitRef="u">abc</t:Revenue></xbrl>'''
        result=self.review('test.xbrl',xml,config={'xbrl_validation':'arelle'})
        self.assertTrue(any(l['rule']=='XBRL_TAXONOMY' and l['status']=='NOT_VERIFIABLE' for l in result['logs']))
        self.assertTrue(any(f['rule_id']=='XBRL_VALIDATION' for f in result['findings']))

if __name__=='__main__':unittest.main()

class ArelleIntegrationTests(unittest.TestCase):
    @unittest.skipUnless(importlib.util.find_spec('arelle'),'Arelle dependency not installed')
    def test_calculation_linkbase_and_clearing(self):
        import shutil
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            source=Path(__file__).resolve().parents[1]/'examples'/'xbrl'
            for p in source.iterdir():
                if p.suffix in ('.xbrl','.xml','.xsd'):
                    shutil.copy(p,root/p.name)
            first=inspect_report_package(str(root/'report.xbrl'),review_config={'require_references':False},output_path=str(root/'first.xlsx'))
            self.assertTrue(any(f['actual']=='xbrl.5.2.5.2:calcInconsistency' for f in first['findings']))
            fixed=root/'fixed.xbrl';fixed.write_text((root/'report.xbrl').read_text().replace('>35</t:Assets>','>30</t:Assets>'))
            second=recheck_corrected_report(first['output_file'],str(fixed),output_path=str(root/'second.xlsx'))
            self.assertEqual(second['summary']['RESOLVED'],1)
            self.assertEqual(second['final_status'],'REVIEW_COMPLETE')

class ExtractionTests(unittest.TestCase):
    def test_image_only_pdf_not_passed(self):
        from pypdf import PdfWriter
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'blank.pdf';writer=PdfWriter();writer.add_blank_page(width=200,height=200)
            with p.open('wb') as out:writer.write(out)
            result=inspect_report_package(str(p),export_excel=False)
            self.assertEqual(result['final_status'],'REVIEW_REQUIRED')
            self.assertTrue(any(l['rule']=='PDF_TEXT' and l['status']=='NOT_VERIFIABLE' for l in result['logs']))
    def test_bad_calculation_shape_returns_failure(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'p.json';p.write_text('{"facts":[],"calculations":[1]}')
            self.assertEqual(inspect_report_package(str(p),export_excel=False)['final_status'],'VALIDATION_FAILED')

class PDFTableTests(unittest.TestCase):
    def test_real_pdf_table_extraction(self):
        path=Path(__file__).resolve().parents[1]/'examples'/'tables'/'report.pdf'
        result=inspect_report_package(str(path),export_excel=False,review_config={'metadata':{'entity':'ABC','scope':'separate','unit':'KRW'},'require_references':False})
        self.assertEqual(len(result['facts']),3)
        finding=next(f for f in result['findings'] if f['rule_id']=='FOOTING')
        self.assertEqual(finding['difference'],'5')
        self.assertTrue(finding['section'].startswith('page:1:table:1'))

class ArchiveXBRLTests(unittest.TestCase):
    @unittest.skipUnless(importlib.util.find_spec('arelle'),'Arelle dependency not installed')
    def test_zip_taxonomy_calculation_validation(self):
        source=Path(__file__).resolve().parents[1]/'examples'/'xbrl'
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'report.zip'
            with ZipFile(p,'w') as z:
                for f in source.iterdir():
                    if f.suffix in ('.xbrl','.xml','.xsd'):
                        z.write(f,'package/'+f.name)
            result=inspect_report_package(str(p),review_config={'require_references':False},export_excel=False)
            self.assertTrue(any(f['rule_id']=='XBRL_VALIDATION' for f in result['findings']))
            self.assertFalse(any(l['status']=='FAILED' for l in result['logs']))
            self.assertFalse(any('frr-archive-' in p for p in result['input_hashes']))

class ClearingSafetyTests(unittest.TestCase):
    def test_deleting_typo_paragraph_is_not_verified_correction(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);p=root/'old.txt';p.write_text('유형 자산은 취득원가로 기록합니다.\n기타 문단')
            first=inspect_report_package(str(p),output_path=str(root/'first.xlsx'))
            corrected=root/'new.txt';corrected.write_text('기타 문단')
            second=recheck_corrected_report(first['output_file'],str(corrected),output_path=str(root/'second.xlsx'))
            self.assertEqual(second['summary']['RESOLVED'],0)
            self.assertEqual(second['summary']['NOT_VERIFIABLE'],1)
    def test_dropped_bad_company_paragraph_not_silently_cleared(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);p=root/'old.txt';p.write_text('회사명: WRONG\n회사명: ABC')
            first=inspect_report_package(str(p),review_config={'expectations':[{'id':'company','label':'회사','selector':'회사명:','expected':'ABC'}]},output_path=str(root/'first.xlsx'))
            corrected=root/'new.txt';corrected.write_text('회사명: ABC')
            second=recheck_corrected_report(first['output_file'],str(corrected),output_path=str(root/'second.xlsx'))
            self.assertEqual(second['summary']['NOT_VERIFIABLE'],1)

class InlineXBRLTests(unittest.TestCase):
    @unittest.skipUnless(importlib.util.find_spec('arelle'),'Arelle dependency not installed')
    def test_inline_facts_and_calculation(self):
        path=Path(__file__).resolve().parents[1]/'examples'/'xbrl'/'inline.xhtml'
        result=inspect_report_package(str(path),export_excel=False,review_config={'require_references':False})
        self.assertEqual(len(result['facts']),3)
        self.assertEqual(result['final_status'],'REVIEW_COMPLETE_WITH_OPEN_ITEMS')
        self.assertTrue(any(f['actual']=='xbrl.5.2.5.2:calcInconsistency' for f in result['findings']))
