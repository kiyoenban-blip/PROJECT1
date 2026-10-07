import json
import hashlib
from pathlib import Path
from uuid import uuid4
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

STATUSES = ['미확인','수정필요','수정중','수정완료','수정불필요','추가확인','Client 확인','Manager 확인','Partner 확인','Closed']
HEADERS = ['No','Issue ID','중요도','Risk','파일','위치','구분','현재내용','비교대상','비교값','차이','중요성대비','문제점','수정제안','근거','Confidence','상태','담당자','검토자','메모']

def safe_cell(ws, row, col, value):
    if value is None:
        return
    if isinstance(value, (dict,list)):
        value = json.dumps(value,ensure_ascii=False)
    if isinstance(value,str):
        # Report content is untrusted. Always write literal strings, never formulas.
        value = ''.join(c for c in value if ord(c)>=32 or c in '\n\t\r')
        if len(value)>32767:
            value=value[:32720]+' [표시 생략: 전체 내용은 _ReviewData에 보존]'
    c = ws.cell(row,col,value)
    if isinstance(value,str):
        c.data_type='s'

def table(wb, name, headers, rows):
    ws=wb.create_sheet(name)
    for r, row in enumerate([headers]+list(rows),1):
        for c,value in enumerate(row,1):
            safe_cell(ws,r,c,value)
    ws.freeze_panes='A2'
    ws.auto_filter.ref=ws.dimensions
    for c in ws[1]:
        c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='17365D')
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width=24
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment=Alignment(vertical='top',wrap_text=True)
    return ws

def dict_table(wb,name,records):
    if not records:
        return
    keys=list(dict.fromkeys(k for item in records for k in item))
    table(wb,name,keys,[[item.get(k) for k in keys] for item in records])

def write_review(result, output_path):
    primary=Path(result['primary_file'])
    path=Path(output_path) if output_path else Path.cwd()/'output'/f'{primary.stem}_Review_{uuid4().hex[:8]}.xlsx'
    path=path.resolve()
    if path.suffix.lower()!='.xlsx':
        raise ValueError('출력 확장자는 .xlsx여야 합니다.')
    protected=set(result['input_hashes']) | {result['primary_file']} | set(result['reference_files']) | {result.get('prior_file'),result.get('previous_version')}
    if str(path) in protected:
        raise ValueError('입력 파일을 출력 경로로 사용할 수 없습니다.')
    if path.exists():
        raise FileExistsError('기존 파일을 덮어쓰지 않습니다. 새 출력 경로를 지정하십시오.')
    wb=Workbook(); wb.remove(wb.active)
    table(wb,'01_검토요약',['항목','값'],[
        ['최종상태',result['final_status']],['검사범위 집계',result.get('coverage',{})],['의미','정의된 검사 범위의 실행 결과이며 감사의견이나 적정성을 보증하지 않습니다.'],
        ['원본','읽기 전용: 자동수정 없음'],['검사버전',result.get('engine_version','0.6.0')],['Confidence','규칙 기반 휴리스틱 점수이며 통계적으로 보정된 확률이 아닙니다.'],
        *[[k,v] for k,v in result['summary'].items()],['기존 Issue',result.get('previous_issue_count','해당 없음')]])
    rows=[]
    for i,f in enumerate(result['findings'],1):
        rows.append([i,f['issue_id'],f['severity'],{'CRITICAL':100,'HIGH':75,'MEDIUM':50,'LOW':25}[f['severity']],
            f['primary_file'], f.get('section') or f.get('cell'),f['category'],f.get('actual') or f.get('original_text'),
            f.get('reference_file'),f.get('expected'),f.get('difference'),
            float(f['materiality_ratio']) if f.get('materiality_ratio') is not None else None,
            f['message'],f.get('suggestion'), '\n'.join(f['evidence']),f['confidence'],f['status'],f['owner'],f['reviewer'],f['memo']])
    ws=table(wb,'02_수정사항',HEADERS,rows)
    dv=DataValidation(type='list',formula1='"'+','.join(STATUSES)+'"'); ws.add_data_validation(dv); dv.add(f'Q2:Q{max(2,len(rows)+1)}')
    for row in ws.iter_rows(min_row=2):
        row[11].number_format='0.0%'; row[15].number_format='0%'
    dict_table(wb,'03_통합숫자대사',result['lineage'])
    dict_table(wb,'04_DataLineage',result['facts'])
    dict_table(wb,'05_문서정합성',[{'Issue ID':f['issue_id'],'clearing':f['clearing'],'rule':f['rule_id'],
        'PM 대비':f.get('performance_materiality_ratio'),'Trivial 대비':f.get('trivial_threshold_ratio'),
        'recommendation':f.get('recommendation_type'),'재검사 근거':f.get('verification'),
        '제안문구':f.get('suggested_text'),'단위':f.get('unit'),'문제점':f['message'],'근거':f['evidence']} for f in result['findings']])
    dict_table(wb,'06_전기비교',result['prior_changes'])
    dict_table(wb,'07_버전비교',result['version_changes'])
    for title, cat in [('08_Excel수식검사','EXCEL'),('09_XBRL검토','STRUCTURE')]:
        dict_table(wb,title,[{'Issue ID':f['issue_id'],'위치':f.get('section'),'문제점':f['message']} for f in result['findings'] if f['category']==cat])
    dict_table(wb,'10_PASS항목',result['passes'])
    dict_table(wb,'11_검사로그',result['logs'])
    table(wb,'12_규칙설정',['규칙','값'], [['schema',result['schema']],['Risk','CRITICAL=100, HIGH=75, MEDIUM=50, LOW=25; 우선순위 표시값'],
        ['Threshold unit',result.get('materiality_unit')],['검사설정',result.get('review_config',{})],
        ['Confidence 등급','>=95% 높은 확률로 수정 필요; >=80% 검토 필요; >=60% 추가 확인 필요; <60% 참고사항'],
        ['입력 SHA256',result['input_hashes']], *[[k,v] for k,v in result['thresholds'].items()]])
    meta=wb.create_sheet('_ReviewData'); meta.sheet_state='veryHidden'
    payload=json.dumps(result,ensure_ascii=True)
    safe_cell(meta,1,2,hashlib.sha256(payload.encode()).hexdigest())
    for i,start in enumerate(range(0,len(payload),30000),1):
        safe_cell(meta,i,1,payload[start:start+30000])
    path.parent.mkdir(parents=True,exist_ok=True)
    # Exclusive creation also prevents accidental overwrite through hardlinks/symlinks.
    created=False
    try:
        with path.open('xb') as stream:
            created=True
            wb.save(stream)
    except FileExistsError:
        raise
    except Exception:
        if created:
            path.unlink(missing_ok=True)
        raise
    finally:
        wb.close()
    return str(path)

def read_review(path):
    wb=load_workbook(path,data_only=False,read_only=True)
    try:
        if '_ReviewData' not in wb or '02_수정사항' not in wb:
            raise ValueError('FRR Review Workbook이 아닙니다.')
        payload=''.join(str(row[0]) for row in wb['_ReviewData'].iter_rows(values_only=True))
        checksum=wb['_ReviewData']['B1'].value
        if checksum and hashlib.sha256(payload.encode()).hexdigest()!=checksum:
            raise ValueError('숨김 검토 데이터의 무결성 검증 실패')
        result=json.loads(payload)
        if result.get('schema')!='frr-review-0.6':
            raise ValueError('지원하지 않는 Review schema')
        records={f['issue_id']:f for f in result['findings']}
        if len(records)!=len(result['findings']):
            raise ValueError('중복 Issue ID')
        seen=set()
        for row in wb['02_수정사항'].iter_rows(min_row=2,values_only=True):
            if not row or row[1] is None:
                continue
            issue=row[1]
            if issue not in records or issue in seen:
                raise ValueError('수정사항 시트의 Issue ID가 변경되었거나 중복되었습니다.')
            seen.add(issue)
            for index,key in [(16,'status'),(17,'owner'),(18,'reviewer'),(19,'memo')]:
                if index>=len(row):
                    raise ValueError('필수 검토 열이 삭제되었습니다.')
                records[issue][key]=str(row[index] or '')
            if records[issue]['status'] not in STATUSES:
                raise ValueError('유효하지 않은 검토 상태')
        if seen!=set(records):
            raise ValueError('수정사항 행이 삭제되었습니다. 상태/담당자/검토자/메모만 편집하십시오.')
        return result
    finally:
        wb.close()
