import csv
import io
import hashlib
import json
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path
from lxml import etree, html
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from pypdf import PdfReader
from .models import Document, Fact, Finding

ERRORS = {'#REF!', '#VALUE!', '#DIV/0!', '#NAME?', '#NUM!', '#N/A', '#NULL!', '#SPILL!', '#CALC!'}
HEADERS = {'concept', 'value', 'period', 'unit', 'entity', 'scope'}

def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def number(value):
    if isinstance(value, bool) or value is None:
        raise ValueError('숫자가 아닙니다')
    s = str(value).strip().replace(',', '')
    if s.startswith('(') and s.endswith(')'):
        s = '-' + s[1:-1]
    n = Decimal(s)
    if not n.is_finite():
        raise ValueError('유한한 숫자가 아닙니다')
    return n

def finding(doc, rule, location, message, actual, expected=None, severity='MEDIUM', **kwargs):
    return Finding(rule_id=rule, category=kwargs.pop('category', 'EXCEL'), severity=severity,
        primary_file=doc.path, section=location, actual=str(actual), expected=None if expected is None else str(expected),
        message=message, evidence=[f'{doc.path} [{location}] {actual}'] + ([f'비교 근거: {expected}'] if expected is not None else []),
        suggestion=kwargs.pop('suggestion', '사용자 확인 후 수정 여부를 결정하십시오.'),
        anchor=f'{rule}|{location}', **kwargs)

def add_rows(doc, rows, prefix):
    for i, row in enumerate(rows, 2):
        location = str(row.get('location') or f'{prefix}{i}')
        try:
            if not HEADERS.issubset(row) or any(row[k] is None or str(row[k]).strip()=='' for k in HEADERS):
                raise ValueError('필수 메타데이터 누락')
            doc.facts.append(Fact(concept=str(row['concept']), value=number(row['value']), period=str(row['period']),
                unit=str(row['unit']), entity=str(row['entity']), scope=str(row['scope']),
                dimensions=str(row.get('dimensions') or ''), location=location, source=doc.path))
        except (ValueError, InvalidOperation) as e:
            doc.logs.append({'rule':'FACT_PARSE','status':'NOT_VERIFIABLE','location':location,'detail':str(e)})

def excel(doc, config):
    wb = load_workbook(doc.path, data_only=False, read_only=False, keep_links=True)
    cached = load_workbook(doc.path, data_only=True, read_only=False)
    try:
        for ws in wb:
            if ws.max_row * ws.max_column > config.max_cells:
                raise ValueError('Excel 셀 수 제한 초과')
            for row in list(ws.rows):
                for c in row:
                    if c.value is None:
                        continue
                    loc = f'{ws.title}!{c.coordinate}'
                    val = str(c.value)
                    doc.texts[loc] = val
                    doc.checked.append(f'FORMULA_ERROR|{loc}')
                    doc.checked.append(f'EXTERNAL_LINK|{loc}')
                    cv = cached[ws.title][c.coordinate]
                    if c.data_type == 'e' or cv.data_type == 'e' or val in ERRORS or (c.data_type == 'f' and '#REF!' in val):
                        doc.findings.append(finding(doc, 'FORMULA_ERROR', loc, 'Excel 수식 오류', cv.value if cv.data_type == 'e' else val,
                            severity='HIGH', sheet=ws.title, cell=c.coordinate, confidence=1))
                    if c.data_type == 'f':
                        if re.search(r'\[[^\]]+\][^!]*!', val):
                            doc.findings.append(finding(doc, 'EXTERNAL_LINK', loc, '외부 파일 참조: 당기 자료인지 확인하십시오.', val,
                                severity='HIGH', sheet=ws.title, cell=c.coordinate, confidence=1))
                        if cv.value is None:
                            if not any(f.rule_id == 'FORMULA_ERROR' and f.section == loc for f in doc.findings):
                                doc.checked.remove(f'FORMULA_ERROR|{loc}')
                            doc.logs.append({'rule':'FORMULA_EVALUATION','status':'NOT_VERIFIABLE','location':loc,'detail':'수식 캐시 없음. 수식을 실행하거나 재계산하지 않습니다.'})
                    if c.data_type == 'f':
                        doc.checked.append(f'HARD_CODE|{loc}')
                    if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                        for dr, dc in [(0,1),(1,0)]:
                            if c.row-dr < 1 or c.column-dc < 1:
                                continue
                            a = ws.cell(c.row-dr, c.column-dc)
                            b = ws.cell(c.row+dr, c.column+dc)
                            if a.data_type == b.data_type == 'f':
                                try:
                                    left = Translator(a.value, origin=a.coordinate).translate_formula(c.coordinate)
                                    right = Translator(b.value, origin=b.coordinate).translate_formula(c.coordinate)
                                    if left == right:
                                        doc.checked.append(f'HARD_CODE|{loc}')
                                        doc.findings.append(finding(doc, 'HARD_CODE', loc, '수식 패턴 사이에 직접 입력된 값이 있습니다.', val, left,
                                            sheet=ws.title, cell=c.coordinate, confidence=.85, suggested_text=left))
                                        break
                                except (ValueError, TypeError, IndexError):
                                    pass
            headers = [c.value for c in ws[1]]
            if HEADERS.issubset(set(headers)):
                rows = []
                for row in cached[ws.title].iter_rows(min_row=2):
                    if all(c.value is None for c in row):
                        continue
                    item = dict(zip(headers, [c.value for c in row]))
                    item['location'] = f'{ws.title}!{row[headers.index("value")].coordinate}'
                    rows.append(item)
                add_rows(doc, rows, '')
            else:
                doc.tables.append({'location':ws.title, 'rows':[[c.value for c in row] for row in cached[ws.title]],
                    'cells':{f'{c.row-1},{c.column-1}':f'{ws.title}!{c.coordinate}' for row in ws for c in row}})
        doc.logs.append({'rule':'EXCEL_STATIC','status':'PASS','detail':'정적 수식/외부링크/Hard Coding 검사 완료. Excel 계산 엔진 미실행.'})
    finally:
        wb.close()
        cached.close()

def read_bytes(doc, data, ext, config):
    from .documents import markup, xml_root
    from .xbrl import read_xbrl
    doc.kind=ext
    if ext == '.csv':
        add_rows(doc, csv.DictReader(io.StringIO(data.decode('utf-8-sig'))), 'row:')
    elif ext == '.json':
        payload=json.loads(data.decode('utf-8-sig'))
        if isinstance(payload,dict):
            doc.metadata.update(payload.get('metadata', {}))
            calcs=payload.get('calculations', [])
            if not isinstance(calcs,list) or any(not isinstance(c,dict) for c in calcs):
                raise ValueError('calculations는 객체 배열이어야 합니다.')
            doc.calculations.extend(calcs)
            texts=payload.get('texts', {})
            if not isinstance(texts,dict) or any(not isinstance(v,str) for v in texts.values()):
                raise ValueError('texts는 문자열 값의 객체여야 합니다.')
            doc.texts.update(texts)
            doc.tables.extend(payload.get('tables', []))
            add_rows(doc,payload.get('facts', []),'fact:')
        else:
            add_rows(doc,payload,'fact:')
    elif ext in ('.xml','.xbrl'):
        root=xml_root(data)
        if etree.QName(root).localname=='xbrl':
            read_xbrl(doc,root,config)
            doc.kind='xbrl'
        else:
            markup(doc,data,xml=True)
    elif ext in ('.html','.htm','.xhtml'):
        root=markup(doc,data)
        if any('nonfraction' in str(e.tag).lower() for e in root.iter()):
            doc.logs.append({'rule':'INLINE_XBRL','status':'NOT_VERIFIABLE','detail':'Inline XBRL에는 Arelle 검증을 사용하십시오. HTML 표 추출은 별도로 수행됩니다.'})
            doc.kind='ixbrl'
    elif ext in ('.txt','.md'):
        doc.texts={f'line:{i}':t for i,t in enumerate(data.decode('utf-8-sig').splitlines(),1) if t.strip()}
    else:
        raise ValueError(f'지원하지 않는 형식: {ext}')


def parse(path, config=None):
    from .config import ReviewConfig
    from .documents import archive
    from .tables import extract_tables, normalize_facts
    config=config or ReviewConfig()
    p=Path(path).resolve(strict=True)
    if not p.is_file() or p.stat().st_size > config.max_file_bytes:
        raise ValueError('입력 파일이 아니거나 파일 크기 제한을 초과했습니다.')
    metadata=dict(config.metadata)
    metadata.update(config.document_metadata.get(str(p), config.document_metadata.get(p.name,{})))
    doc=Document(path=str(p),sha256=digest(p),metadata=metadata,kind=p.suffix.lower())
    ext=p.suffix.lower()
    if ext in ('.xlsx','.xlsm'):
        excel(doc,config)
    elif ext in ('.zip','.dsd'):
        archive(doc,p.read_bytes(),config,read_bytes)
    elif ext=='.pdf':
        import pdfplumber
        with pdfplumber.open(p) as pdf:
            for i,page in enumerate(pdf.pages,1):
                text=page.extract_text() or ''
                for j,line in enumerate(text.splitlines(),1):
                    doc.texts[f'page:{i}:line:{j}']=line
                if not text.strip():
                    doc.logs.append({'rule':'PDF_TEXT','status':'NOT_VERIFIABLE','location':f'page:{i}','detail':'텍스트 없는 페이지: OCR이 필요합니다.'})
                for j,rows in enumerate(page.extract_tables(),1):
                    doc.tables.append({'location':f'page:{i}:table:{j}','rows':rows})
    else:
        read_bytes(doc,p.read_bytes(),ext,config)
    extract_tables(doc,config)
    normalize_facts(doc,config)
    if not doc.facts:
        doc.logs.append({'rule':'FACT_EXTRACTION','status':'NOT_VERIFIABLE','detail':'정규 Fact가 없습니다. 메타데이터/표 구조 또는 OCR 결과가 필요합니다.'})
    return doc
